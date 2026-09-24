#!/usr/bin/env python3
"""
生成 PPT 导入用的**第三方写入器**夹具（第 99 轮）。

为什么用 python-pptx 而不是我们自己的导出器：O-12 ① 要的是"打开一份**真实** deck 看看渲染"
（`.pptx` 一直没走查，因为本机没有可打开的演示文稿）。用**别的写入器**产出的文件才有意义 ——
自己写、自己读等于自己跟自己对照。python-pptx 生成的是标准 OOXML，PowerPoint 能打开。

夹具内容（每页都刻意放**不同形态**的元素，用来量导入器的覆盖面）：
  1. 标题页（Title Slide 版式：标题 + 副标题）
  2. 标题 + 项目符号（三条要点）—— 正文占位符
  3. 标题 + **表格**（2×2）—— 表格是 `p:graphicFrame` 而不是 `p:sp`
  4. 标题 + **图片**（内联 PNG）—— 走 `p:pic` + `a:blip` 的媒体提取
  5. 标题 + **演讲者备注** —— 备注在 `ppt/notesSlides/`

跑法（本机没有把 python-pptx 装进项目依赖，所以这是一次性生成脚本，夹具本身已入库）：
  "<python-pptx 所在环境的 python>" tools/fixtures/make-pptx-fixture.py src/test/fixtures/pptx/third-party-deck.pptx
"""
import struct
import sys
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches


def tiny_png(width: int = 16, height: int = 16, rgb=(220, 60, 60)) -> bytes:
    """不依赖 Pillow 的纯字节 PNG（纯色块）"""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def main(out_path: str) -> None:
    dest = Path(out_path)
    dest.parent.mkdir(parents=True, exist_ok=True)  # 临时 PNG 也要先有目录
    prs = Presentation()
    # 16:9 ⇒ 1280×720 px（导入器按 EMU→px 换算，正好能被断言钉住）
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ── 1 标题页 ──
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Codem PPTX 导入钻取"
    s1.placeholders[1].text = "由 python-pptx 生成（第三方写入器）"

    # ── 2 项目符号 ──
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "三条要点"
    tf = s2.placeholders[1].text_frame
    tf.text = "第一条要点"
    for extra in ("第二条要点", "第三条要点"):
        para = tf.add_paragraph()
        para.text = extra

    # ── 3 表格 ──
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "表格页"
    table = s3.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(6), Inches(2)
    ).table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "通过率"
    table.cell(1, 1).text = "100%"

    # ── 4 图片 ──
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "图片页"
    tmp_png = Path(out_path).with_suffix(".png")
    tmp_png.write_bytes(tiny_png())
    s4.shapes.add_picture(str(tmp_png), Inches(1), Inches(3), Inches(1), Inches(1))
    tmp_png.unlink()

    # ── 5 备注 ──
    s5 = prs.slides.add_slide(prs.slide_layouts[5])
    s5.shapes.title.text = "备注页"
    s5.notes_slide.notes_text_frame.text = "这段是演讲者备注（量导入器读不读它）"

    prs.core_properties.title = "第三方 PPTX 夹具"
    prs.core_properties.author = "codem audit fixture"

    prs.save(str(dest))
    print(f"已写入 {dest}（{dest.stat().st_size} 字节）")
    print("夹具保持 python-pptx 的原始输出（未做任何改写）：单引号 XML 声明对 jsdom / 真实 Chromium 都合法。")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "src/test/fixtures/pptx/third-party-deck.pptx"
    main(target)
